"""Sincronización mínima para visor PÚBLICO: SOLO ID inmueble y coordenadas.
No escribir nunca la base completa en el repositorio ni mostrar sus campos privados.
"""
import json, os, math, urllib.parse, urllib.request, urllib.error
from io import BytesIO
from pathlib import Path
from openpyxl import load_workbook


def fetch(req, label, binary=False):
    try:
        with urllib.request.urlopen(req, timeout=120) as res:
            return res.read() if binary else json.load(res)
    except urllib.error.HTTPError as exc:
        message = exc.read().decode('utf-8', errors='replace')
        try:
            data = json.loads(message)
            error = data.get('error', {})
            message = error.get('message', message) if isinstance(error, dict) else data.get('error_description', message)
        except ValueError:
            pass
        raise SystemExit(f'{label}: HTTP {exc.code}. {str(message)[:300]}') from None


def run():
    names = ('AZURE_CLIENT_ID','AZURE_TENANT_ID','ONEDRIVE_DRIVE_ID','ONEDRIVE_FILE_ID')
    absent = [n for n in names if not os.getenv(n, '').strip()]
    if absent: raise SystemExit('Faltan variables: ' + ', '.join(absent))
    endpoint = os.environ['ACTIONS_ID_TOKEN_REQUEST_URL']
    endpoint += ('&' if '?' in endpoint else '?') + 'audience=' + urllib.parse.quote('api://AzureADTokenExchange',safe='')
    assertion = fetch(urllib.request.Request(endpoint,headers={'Authorization': 'Bearer '+os.environ['ACTIONS_ID_TOKEN_REQUEST_TOKEN']}),'GitHub OIDC')['value']
    tenant = os.environ['AZURE_TENANT_ID'].strip()
    token_url = f'https://login.microsoftonline.com/{tenant}/oauth2/v2.0/token'
    params = {'client_id':os.environ['AZURE_CLIENT_ID'].strip(),'scope':'https://graph.microsoft.com/.default','grant_type':'client_credentials','client_assertion_type':'urn:ietf:params:oauth:client-assertion-type:jwt-bearer','client_assertion':assertion}
    req = urllib.request.Request(token_url,urllib.parse.urlencode(params).encode(),headers={'Content-Type':'application/x-www-form-urlencoded'},method='POST')
    access_token=fetch(req,'Microsoft Entra')['access_token']
    drive = urllib.parse.quote(os.environ['ONEDRIVE_DRIVE_ID'].strip(),safe='')
    item = urllib.parse.quote(os.environ['ONEDRIVE_FILE_ID'].strip(),safe='')
    req = urllib.request.Request(f'https://graph.microsoft.com/v1.0/drives/{drive}/items/{item}/content',headers={'Authorization':'Bearer '+access_token})
    raw=fetch(req,'Microsoft Graph / Excel',binary=True)
    workbook=load_workbook(BytesIO(raw),read_only=True,data_only=True)
    if 'BASE_HUB' not in workbook.sheetnames: raise SystemExit('No se encontró BASE_HUB')
    sheet=workbook['BASE_HUB']
    rows=sheet.iter_rows(values_only=True)
    headers=[str(v or '').strip().upper() for v in next(rows)]
    cols={h:i for i,h in enumerate(headers) if h}
    required=['ID_INMUEBLE','LATITUD','LONGITUD']
    for key in required:
        if key not in cols:raise SystemExit(f'Falta columna {key} en BASE_HUB')
    result=[]
    seen=set()
    for excel_row,values in enumerate(rows,start=2):
        if not any(v is not None and str(v).strip() for v in values):continue
        code=str(values[cols['ID_INMUEBLE']] or '').strip()
        # Sin identificador: mostrar referencia de fila, etiquetada explícitamente como tal.
        if not code or code in ('-', '.') or ';' in code:
            code=f'FILA-{excel_row}'
        # No agrupar varias filas con el mismo identificador; asegurar unicidad de selección.
        key=code if code not in seen else f'{code}-FILA-{excel_row}'
        seen.add(key)
        def coordinate(k,lo,hi):
            value=values[cols[k]]
            try:
                n=float(str(value).strip().replace(',','.'))
                return round(n,7) if math.isfinite(n) and lo<=n<=hi else None
            except (TypeError,ValueError):return None
        lat=coordinate('LATITUD',-90,90)
        lon=coordinate('LONGITUD',-180,180)
        # No usar (0,0) como coordenada fiable.
        if lat==0 and lon==0:lat=lon=None
        result.append({'codigo':key,'latitud':lat,'longitud':lon})
    if not result:raise SystemExit('No hay registros; se conserva la versión anterior')
    target=Path('data/inmuebles.json')
    target.parent.mkdir(exist_ok=True,parents=True)
    target.write_text(json.dumps({'inmuebles':result},ensure_ascii=False,separators=(',',':')),encoding='utf-8')
    print('Fuente de mapa actualizada; no se publicaron direcciones, contratos, nombres ni datos financieros.')
    summary=os.environ.get('GITHUB_STEP_SUMMARY')
    if summary:
        with open(summary,'a',encoding='utf-8') as f:
            f.write('### Actualización de ubicación de inmuebles\n\nArchivo de datos mínimo generado correctamente. Sin indicadores públicos. Sin información contractual o personal.\n')

if __name__=='__main__':run()
